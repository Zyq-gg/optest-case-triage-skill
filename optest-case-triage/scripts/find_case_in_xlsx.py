#!/usr/bin/env python3
"""Find PyTorch optest cases in an Excel workbook.

Examples:
  python3 find_case_in_xlsx.py --xlsx /workspace/pytorch2.12.0-optest_2_marked_newcases.xlsx --op-name test_max_min_bool_cpu
  python3 find_case_in_xlsx.py --xlsx report.xlsx --py-name test/test_ops.py --class-name TestCommonCPU --op-name test_max_min_bool_cpu --exact
"""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from openpyxl import load_workbook


def clean(value) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def row_values(ws, row: int) -> dict[str, str]:
    headers = [clean(ws.cell(1, col).value) or f"col_{col}" for col in range(1, ws.max_column + 1)]
    values = {}
    for col, header in enumerate(headers, start=1):
        value = ws.cell(row, col).value
        if value not in (None, ""):
            values[header] = clean(value)
    return values


def matches(value: str, query: str, exact: bool) -> bool:
    if not query:
        return True
    value = clean(value)
    query = clean(query)
    return value == query if exact else query in value


def main() -> int:
    parser = argparse.ArgumentParser(description="Find optest case rows in an Excel workbook.")
    parser.add_argument("--xlsx", type=Path, required=True)
    parser.add_argument("--py-name", default="")
    parser.add_argument("--class-name", default="")
    parser.add_argument("--op-name", default="", required=True)
    parser.add_argument("--sheet", default="", help="Optional sheet name filter")
    parser.add_argument("--exact", action="store_true", help="Require exact matches for provided fields")
    parser.add_argument("--json", action="store_true", help="Emit JSON instead of readable text")
    args = parser.parse_args()

    wb = load_workbook(args.xlsx, data_only=True, read_only=False)
    results = []
    for ws in wb.worksheets:
        if args.sheet and ws.title != args.sheet:
            continue
        for row in range(2, ws.max_row + 1):
            py_name = clean(ws.cell(row, 1).value)
            class_name = clean(ws.cell(row, 2).value)
            op_name = clean(ws.cell(row, 3).value)
            if not op_name:
                continue
            if not matches(py_name, args.py_name, args.exact):
                continue
            if not matches(class_name, args.class_name, args.exact):
                continue
            if not matches(op_name, args.op_name, args.exact):
                continue
            results.append(
                {
                    "sheet": ws.title,
                    "row": row,
                    "py_name": py_name,
                    "class_name": class_name,
                    "op_name": op_name,
                    "values": row_values(ws, row),
                }
            )
    wb.close()

    if args.json:
        print(json.dumps(results, ensure_ascii=False, indent=2))
        return 0

    print(f"matches={len(results)}")
    for item in results:
        print(f"\n[{item['sheet']}!{item['row']}] {item['py_name']}::{item['class_name']}::{item['op_name']}")
        for key, value in item["values"].items():
            print(f"  {key}: {value}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
